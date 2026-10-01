"""월마감 반품 집계: 임시 SQLite DB와 다운로드 검증."""
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
import sys
import gc
import tempfile
import unittest
import pandas as pd
from openpyxl import load_workbook
ROOT = Path(__file__).resolve().parents[1] / "order_management_src"
for folder in (ROOT, ROOT / "app_layers"):
    sys.path.insert(0, str(folder))
from app_layers.db_migration import initialize_database
from repositories import purchase_repository as repo
from services.purchase_service import PurchaseService
from services.statement_returns import return_marker
from ui.pages import accounting_export as accounting

class AccountingReturnsTests(unittest.TestCase):
    def test_temporary_database_summary_and_excel(self):
        with tempfile.TemporaryDirectory() as folder:
            folder = Path(folder)
            initialize_database(folder)
            headers = []
            items = []
            # 순상품금액, 반품금액, 배송비: 부분반품/전액반품/반품없음/다른월
            for sid, net, returned, freight, day in [
                ("S1", 7000, 3000, 500, "2026-09-01"),
                ("S2", 0, 9501, 200, "2026-09-02"),
                ("S3", 4000, 0, 0, "2026-09-03"),
                ("S4", 1000, 1000, 0, "2026-08-01")]:
                header = dict.fromkeys(repo.STATEMENT_COLUMNS, "")
                header.update(명세서ID=sid, 발주ID="O1", 거래처명="거래처A", 명세서번호=sid, 명세서일자=day, 운송비=freight)
                headers.append(header)
                item = dict.fromkeys(repo.STATEMENT_ITEM_COLUMNS, "")
                item.update(명세서ID=sid, 순번=1, 제품코드="P1", 입고수량=net//1000, 매입단가=1000, 상품금액=net, 가격적용여부=return_marker(3, returned) if returned else "적용")
                items.append(item)
            repo.replace_statements(folder, pd.DataFrame(headers))
            repo.replace_statement_items(folder, pd.DataFrame(items))
            service = PurchaseService(folder)
            before = service.load_all()
            purchase = SimpleNamespace(load_purchase_data=service.load_all, to_int=repo._int)
            summary, detail, vendor = accounting._build_frames(purchase, {"orders":pd.DataFrame()}, 2026, 9)
            self.assertEqual(list(summary.columns)[-4:], ["상품금액", "배송비", "반품금액", "총 매입금액"])
            self.assertEqual(summary["상품금액"].tolist(), [10000, 9501, 4000])
            self.assertEqual(summary["반품금액"].tolist(), [3000, 9501, 0])
            self.assertEqual(summary["총 매입금액"].tolist(), [7500, 200, 4000])
            self.assertEqual(vendor.iloc[0]["총매입금액"], 11700)
            self.assertEqual(vendor.iloc[0]["반품금액"], 12501)
            self.assertEqual(pd.to_numeric(detail["금액"]).sum(), 11700)
            wb = load_workbook(BytesIO(accounting._excel_bytes("2026-09", summary, detail, vendor)))
            ws = wb["월마감 요약"]
            cols = {cell.value:cell.column for cell in ws[3]}
            self.assertEqual(ws.cell(7, cols["반품금액"]).value, 12501)
            self.assertEqual(ws.cell(7, cols["총 매입금액"]).value, 11700)
            self.assertEqual(ws.cell(4, cols["반품금액"]).number_format, "#,##0")
            for left, right in zip(before, service.load_all()):
                pd.testing.assert_frame_equal(left, right)
            self.assertTrue(accounting._build_frames(purchase, {"orders":pd.DataFrame()}, 2027, 1)[0].empty)
            gc.collect()

if __name__ == "__main__":
    unittest.main()
